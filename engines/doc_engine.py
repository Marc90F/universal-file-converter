"""
Motore Documenti basato su pdf2docx, docx2pdf, PyMuPDF e ReportLab.
Ricostruisce layout, paragrafi, tabelle e flussi vettoriali tra PDF, DOCX e formati correlati.
Include fallback automatici in caso di assenza di Microsoft Word su PC.
"""

import os
from typing import Dict, Any, Callable, Optional, List

# Prova importazione motori documenti
try:
    from pdf2docx import Converter as PdfConverter
    PDF2DOCX_AVAILABLE = True
except ImportError:
    PDF2DOCX_AVAILABLE = False

try:
    from docx2pdf import convert as docx2pdf_convert
    DOCX2PDF_AVAILABLE = True
except ImportError:
    DOCX2PDF_AVAILABLE = False

try:
    import pymupdf  # Modern PyMuPDF API
    PYMUPDF_AVAILABLE = True
except ImportError:
    PYMUPDF_AVAILABLE = False

try:
    from docx import Document
    PYTHON_DOCX_AVAILABLE = True
except ImportError:
    PYTHON_DOCX_AVAILABLE = False

try:
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False


def convert_document(
    src_path: str,
    dst_path: str,
    options: Optional[Dict[str, Any]] = None,
    progress_callback: Optional[Callable[[int, str], None]] = None
) -> Dict[str, Any]:
    """
    Esegue la conversione strutturale del documento.
    """
    options = options or {}
    src_ext = os.path.splitext(src_path)[1].lower().lstrip(".")
    dst_ext = os.path.splitext(dst_path)[1].lower().lstrip(".")

    os.makedirs(os.path.dirname(dst_path), exist_ok=True)

    # 1. PDF ➔ DOCX (Ricostruzione strutturale del layout vettoriale)
    if src_ext == "pdf" and dst_ext == "docx":
        if not PDF2DOCX_AVAILABLE:
            raise RuntimeError("Il modulo 'pdf2docx' non è installato. Esegui l'installazione dei componenti.")

        if progress_callback:
            progress_callback(10, f"Analisi vettoriale e glifi tipografici di {os.path.basename(src_path)}...")

        cv = PdfConverter(src_path)
        try:
            if progress_callback:
                progress_callback(40, "Ricostruzione tabelle, blocchi di testo e gerarchia XML...")
            cv.convert(dst_path, start=0, end=None)
        finally:
            cv.close()

        if progress_callback:
            progress_callback(100, f"Documento Word DOCX generato: {os.path.basename(dst_path)}")

    # 2. DOCX ➔ PDF
    elif src_ext == "docx" and dst_ext == "pdf":
        # 1. Metodo prioritario: Word COM tramite docx2pdf (su Windows/macOS se Word è presente)
        if DOCX2PDF_AVAILABLE:
            try:
                if progress_callback:
                    progress_callback(20, "Interrogazione motore tipografico Word per rendering PDF...")
                docx2pdf_convert(src_path, dst_path)
                converted = True
            except Exception:
                pass

        # 2. Metodo per Linux / macOS / Windows open-source: LibreOffice Headless
        if not converted:
            import shutil
            import subprocess
            soffice = shutil.which("libreoffice") or shutil.which("soffice")
            if soffice:
                try:
                    if progress_callback:
                        progress_callback(25, "Rendering tipografico tramite LibreOffice headless...")
                    out_dir = os.path.dirname(dst_path)
                    subprocess.run(
                        [soffice, "--headless", "--convert-to", "pdf", "--outdir", out_dir, src_path],
                        stdout=subprocess.DEVNULL,
                        stderr=subprocess.DEVNULL,
                        check=True
                    )
                    converted = os.path.exists(dst_path)
                except Exception:
                    pass

        # 3. Fallback puro Python: python-docx + ReportLab (non necessita di alcun Office o software esterno!)
        if not converted:
            if not (PYTHON_DOCX_AVAILABLE and REPORTLAB_AVAILABLE):
                raise RuntimeError("Impossibile convertire DOCX in PDF: Word non trovato e librerie di fallback mancanti.")

            if progress_callback:
                progress_callback(30, "Fallback Python puro: estrazione paragrafi e rendering ReportLab...")
            _render_docx_to_pdf_reportlab(src_path, dst_path, progress_callback)

        if progress_callback:
            progress_callback(100, f"PDF generato con successo: {os.path.basename(dst_path)}")

    # 3. PDF ➔ Immagini (PNG / JPG)
    elif src_ext == "pdf" and dst_ext in ("png", "jpg", "jpeg"):
        if not PYMUPDF_AVAILABLE:
            raise RuntimeError("Il modulo 'pymupdf' è necessario per il rastering ad alta definizione dei PDF.")

        if progress_callback:
            progress_callback(20, "Apertura PDF e rastering pagine a 300 DPI...")

        doc = pymupdf.open(src_path)
        total_pages = len(doc)
        output_files = []

        base, _ = os.path.splitext(dst_path)

        for i, page in enumerate(doc):
            # Render a 300 DPI (72 * 4.16)
            zoom = 300 / 72
            mat = pymupdf.Matrix(zoom, zoom)
            pix = page.get_pixmap(matrix=mat, alpha=False if dst_ext in ("jpg", "jpeg") else True)

            if total_pages == 1:
                page_out = dst_path
            else:
                page_out = f"{base}_pagina_{i+1}.{dst_ext}"

            pix.save(page_out)
            output_files.append(page_out)

            if progress_callback:
                percent = int(20 + ((i + 1) / total_pages) * 75)
                progress_callback(percent, f"Render pagina {i+1} di {total_pages}...")

        doc.close()
        if progress_callback:
            progress_callback(100, f"Estratte {len(output_files)} pagine come immagini.")

    # 4. PDF ➔ Testo / Markdown
    elif src_ext == "pdf" and dst_ext in ("txt", "md"):
        if not PYMUPDF_AVAILABLE:
            raise RuntimeError("Il modulo 'pymupdf' è necessario per l'estrazione testuale dal PDF.")

        if progress_callback:
            progress_callback(30, "Estrazione del flusso testuale e layout...")

        doc = pymupdf.open(src_path)
        text_content = []
        for i, page in enumerate(doc):
            text_content.append(f"--- PAGINA {i+1} ---\n")
            text_content.append(page.get_text())

        doc.close()

        with open(dst_path, "w", encoding="utf-8") as f:
            f.write("\n".join(text_content))

        if progress_callback:
            progress_callback(100, f"Testo estratto con successo: {os.path.basename(dst_path)}")

    # 5. TXT / MD ➔ PDF
    elif src_ext in ("txt", "md") and dst_ext == "pdf":
        if not REPORTLAB_AVAILABLE:
            raise RuntimeError("Il modulo 'reportlab' è necessario per impaginare testo in PDF.")

        if progress_callback:
            progress_callback(30, "Impaginazione tipografica del testo...")

        with open(src_path, "r", encoding="utf-8", errors="replace") as f:
            lines = f.readlines()

        doc = SimpleDocTemplate(dst_path, pagesize=A4, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
        styles = getSampleStyleSheet()
        normal_style = styles['Normal']
        normal_style.fontSize = 10
        normal_style.leading = 14

        story = []
        for line in lines:
            line_str = line.strip().replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            if not line_str:
                story.append(Spacer(1, 10))
            else:
                story.append(Paragraph(line_str, normal_style))

        doc.build(story)

        if progress_callback:
            progress_callback(100, f"PDF tipografico generato: {os.path.basename(dst_path)}")

    else:
        raise ValueError(f"Conversione documentale da '{src_ext}' a '{dst_ext}' non supportata.")

    orig_bytes = os.path.getsize(src_path)
    new_bytes = os.path.getsize(dst_path) if os.path.exists(dst_path) else 0

    return {
        "status": "success",
        "src": src_path,
        "dst": dst_path,
        "orig_bytes": orig_bytes,
        "new_bytes": new_bytes
    }


def _render_docx_to_pdf_reportlab(docx_path: str, pdf_path: str, progress_callback):
    """Fallback per convertire DOCX in PDF quando MS Word non è installato sul computer."""
    doc = Document(docx_path)
    pdf_doc = SimpleDocTemplate(pdf_path, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    story = []

    for p in doc.paragraphs:
        txt = p.text.strip().replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        if not txt:
            story.append(Spacer(1, 8))
            continue

        if p.style.name.startswith("Heading 1"):
            style = styles['Heading1']
        elif p.style.name.startswith("Heading 2"):
            style = styles['Heading2']
        elif p.style.name.startswith("Heading 3"):
            style = styles['Heading3']
        else:
            style = styles['Normal']

        story.append(Paragraph(txt, style))
        story.append(Spacer(1, 4))

    for table in doc.tables:
        table_data = []
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            table_data.append(row_data)

        if table_data:
            t = Table(table_data)
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#1e293b")),
                ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
                ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
                ('BOTTOMPADDING', (0,0), (-1,0), 6),
                ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
            ]))
            story.append(Spacer(1, 8))
            story.append(t)
            story.append(Spacer(1, 8))

    pdf_doc.build(story)
