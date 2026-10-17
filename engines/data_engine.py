"""
Motore Dati e Tabelle basato su Pandas e OpenPyXL (con fallback su moduli standard python).
Analizza la struttura sintattica in un DataFrame in memoria e rigenera la serializzazione
nel formato di destinazione (CSV, XLSX, JSON, HTML, XML).
"""

import os
import csv
import json
import xml.etree.ElementTree as ET
from typing import Dict, Any, Callable, Optional, List

# Prova ad importare pandas e openpyxl
try:
    import pandas as pd
    PANDAS_AVAILABLE = True
except ImportError:
    PANDAS_AVAILABLE = False

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


def _detect_csv_delimiter(file_path: str) -> str:
    """Tenta di dedurre il delimitatore del file CSV (virgola, punto e virgola, tab)."""
    with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
        sample = "".join(f.readline() for _ in range(5))
    try:
        sniffer = csv.Sniffer()
        dialect = sniffer.sniff(sample)
        return dialect.delimiter
    except Exception:
        # Fallback euristico
        if sample.count(";") > sample.count(","):
            return ";"
        elif sample.count("\t") > sample.count(","):
            return "\t"
        return ","


def convert_data(
    src_path: str,
    dst_path: str,
    options: Optional[Dict[str, Any]] = None,
    progress_callback: Optional[Callable[[int, str], None]] = None
) -> Dict[str, Any]:
    """
    Esegue la conversione semantica dei dati attraverso parsing e riserializzazione.
    """
    options = options or {}
    src_ext = os.path.splitext(src_path)[1].lower().lstrip(".")
    dst_ext = os.path.splitext(dst_path)[1].lower().lstrip(".")

    if progress_callback:
        progress_callback(10, f"Apertura e parsing della struttura di {os.path.basename(src_path)}...")

    df = None

    if PANDAS_AVAILABLE:
        # 1. Lettura tramite Pandas
        try:
            if src_ext in ("csv", "tsv"):
                sep = "\t" if src_ext == "tsv" else _detect_csv_delimiter(src_path)
                # Tenta con utf-8, poi fallback su latin1
                try:
                    df = pd.read_csv(src_path, sep=sep, encoding="utf-8")
                except UnicodeDecodeError:
                    df = pd.read_csv(src_path, sep=sep, encoding="latin1")

            elif src_ext in ("xlsx", "xls"):
                df = pd.read_excel(src_path)

            elif src_ext == "json":
                try:
                    df = pd.read_json(src_path)
                except Exception:
                    # Tenta normalizzazione se è un json a oggetti annidati
                    with open(src_path, "r", encoding="utf-8") as f:
                        data = json.load(f)
                    df = pd.json_normalize(data)

            elif src_ext == "html":
                tables = pd.read_html(src_path)
                if not tables:
                    raise ValueError("Nessuna tabella <table> trovata nel documento HTML.")
                df = tables[0]

            elif src_ext == "xml":
                # Utilizza parser='etree' per disabilitare entità esterne e prevenire attacchi XXE
                df = pd.read_xml(src_path, parser="etree")

            else:
                raise ValueError(f"Formato sorgente '{src_ext}' non supportato per i dati.")

        except Exception as e:
            raise RuntimeError(f"Errore durante l'interpretazione del file di dati: {e}")

    else:
        # 2. Fallback con libreria standard (Zero external dependencies)
        if progress_callback:
            progress_callback(20, "Modalità compatibilità Python Standard (senza Pandas)...")
        rows = []
        if src_ext == "csv":
            sep = _detect_csv_delimiter(src_path)
            with open(src_path, 'r', encoding='utf-8', errors='replace') as f:
                reader = csv.DictReader(f, delimiter=sep)
                rows = list(reader)
        elif src_ext == "json":
            with open(src_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
                rows = data if isinstance(data, list) else [data]
        else:
            raise RuntimeError("Per convertire formati complessi come XLSX o XML è necessario installare Pandas.")

        # Conversione standard rows in dizionari
        return _convert_fallback_rows(rows, dst_path, dst_ext, progress_callback)

    if progress_callback:
        progress_callback(50, f"Tabella in memoria: {len(df)} righe x {len(df.columns)} colonne. Rigenerazione...")

    os.makedirs(os.path.dirname(dst_path), exist_ok=True)

    # 3. Scrittura del file target
    if dst_ext == "csv":
        df.to_csv(dst_path, index=False, encoding="utf-8-sig")

    elif dst_ext == "xlsx":
        # Scrittura con formattazione professionale tramite openpyxl
        with pd.ExcelWriter(dst_path, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Dati")
            ws = writer.sheets["Dati"]

            # Stile elegante per riga d'intestazione
            header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

            for col_num, col_name in enumerate(df.columns, 1):
                cell = ws.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

                # Auto-adattamento larghezza colonne
                max_len = max(len(str(val)) for val in [col_name] + list(df[col_name].dropna().head(100)))
                ws.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = min(50, max(max_len + 3, 12))

    elif dst_ext == "json":
        df.to_json(dst_path, orient="records", indent=2, force_ascii=False)

    elif dst_ext == "html":
        # Genera tabella HTML stilizzata e moderna
        html_content = f"""<!DOCTYPE html>
<html lang="it">
<head>
    <meta charset="UTF-8">
    <title>Dati Esportati</title>
    <style>
        body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; margin: 2rem; background: #f8fafc; color: #1e293b; }}
        h2 {{ margin-bottom: 0.5rem; }}
        p.badge {{ display: inline-block; padding: 4px 10px; background: #e2e8f0; border-radius: 6px; font-size: 0.85rem; font-weight: 600; }}
        table {{ border-collapse: collapse; width: 100%; background: #ffffff; box-shadow: 0 4px 6px -1px rgba(0,0,0,0.1); border-radius: 8px; overflow: hidden; }}
        th {{ background: #0f172a; color: #ffffff; text-align: left; padding: 12px 16px; font-size: 0.9rem; text-transform: uppercase; letter-spacing: 0.05em; }}
        td {{ padding: 10px 16px; border-bottom: 1px solid #e2e8f0; font-size: 0.95rem; }}
        tr:nth-child(even) {{ background: #f8fafc; }}
        tr:hover {{ background: #f1f5f9; }}
    </style>
</head>
<body>
    <h2>Tabella Dati Convertita</h2>
    <p class="badge">{len(df)} righe &bull; {len(df.columns)} colonne</p>
    {df.to_html(index=False, classes='table')}
</body>
</html>"""
        with open(dst_path, "w", encoding="utf-8") as f:
            f.write(html_content)

    elif dst_ext == "xml":
        # Genera XML strutturato e pulito
        root = ET.Element("dataset")
        for _, row in df.iterrows():
            record = ET.SubElement(root, "record")
            for col in df.columns:
                col_name_clean = "".join(c if c.isalnum() else "_" for c in str(col)).strip("_")
                if not col_name_clean or col_name_clean[0].isdigit():
                    col_name_clean = f"col_{col_name_clean}"
                field = ET.SubElement(record, col_name_clean)
                val = row[col]
                field.text = "" if pd.isna(val) else str(val)
        tree = ET.ElementTree(root)
        ET.indent(tree, space="  ")
        tree.write(dst_path, encoding="utf-8", xml_declaration=True)

    if progress_callback:
        progress_callback(100, f"Conversione dati completata: {os.path.basename(dst_path)}")

    orig_bytes = os.path.getsize(src_path)
    new_bytes = os.path.getsize(dst_path)

    return {
        "status": "success",
        "src": src_path,
        "dst": dst_path,
        "rows": len(df),
        "columns": len(df.columns),
        "orig_bytes": orig_bytes,
        "new_bytes": new_bytes
    }


def _convert_fallback_rows(rows: List[Dict], dst_path: str, dst_ext: str, progress_callback) -> Dict[str, Any]:
    """Salvataggio di emergenza senza Pandas utilizzando solo moduli nativi."""
    if not rows:
        raise ValueError("Nessun dato trovato da esportare.")

    keys = list(rows[0].keys())

    if dst_ext == "csv":
        with open(dst_path, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.DictWriter(f, fieldnames=keys)
            writer.writeheader()
            writer.writerows(rows)
    elif dst_ext == "json":
        with open(dst_path, 'w', encoding='utf-8') as f:
            json.dump(rows, f, indent=2, ensure_ascii=False)
    elif dst_ext == "xml":
        root = ET.Element("dataset")
        for r in rows:
            record = ET.SubElement(root, "record")
            for k, v in r.items():
                clean_k = "".join(c if c.isalnum() else "_" for c in str(k)).strip("_") or "item"
                field = ET.SubElement(record, clean_k)
                field.text = str(v)
        tree = ET.ElementTree(root)
        ET.indent(tree, space="  ")
        tree.write(dst_path, encoding="utf-8", xml_declaration=True)
    else:
        raise RuntimeError(f"Il formato {dst_ext.upper()} richiede la libreria Pandas per essere generato.")

    if progress_callback:
        progress_callback(100, "Esportazione dati completata!")

    return {"status": "success", "rows": len(rows), "dst": dst_path}
